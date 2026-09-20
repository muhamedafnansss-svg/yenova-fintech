from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from datetime import datetime, timedelta
import io
import csv

from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.database.config import get_db
from app.models.user import User
from app.middleware.deps import get_current_user
from app.models.ledger import Ledger, TransactionType, TransactionStatus
from app.models.project import Project
from app.services.analytics_engine import get_financial_health_score
from app.services.balance import get_current_balance

router = APIRouter(prefix="/api/export", tags=["export"])

def get_report_period(report_type: str):
    """Calculate date range and readable title based on report_type."""
    today = datetime.now()
    if report_type == "daily":
        start_date = today.replace(hour=0, minute=0, second=0, microsecond=0)
        end_date = today.replace(hour=23, minute=59, second=59, microsecond=999999)
        title = "Daily Financial Report"
    elif report_type == "weekly":
        start_date = (today - timedelta(days=today.weekday())).replace(hour=0, minute=0, second=0, microsecond=0)
        end_date = today.replace(hour=23, minute=59, second=59, microsecond=999999)
        title = "Weekly Financial Report"
    elif report_type == "monthly":
        start_date = today.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        end_date = today.replace(hour=23, minute=59, second=59, microsecond=999999)
        title = "Monthly Financial Report"
    elif report_type == "yearly":
        if today.month < 6:
            start_date = today.replace(year=today.year-1, month=6, day=1, hour=0, minute=0, second=0, microsecond=0)
        else:
            start_date = today.replace(month=6, day=1, hour=0, minute=0, second=0, microsecond=0)
        end_date = today.replace(hour=23, minute=59, second=59, microsecond=999999)
        title = "Annual Academic Financial Report"
    else:
        start_date = None
        end_date = None
        title = "General Ledger Report"
    return start_date, end_date, title

@router.get("/pdf")
def export_pdf(
    report_type: str = "daily", 
    project_id: str = None, 
    db: Session = Depends(get_db), 
    current_user: User = Depends(get_current_user)
):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=18)
    styles = getSampleStyleSheet()
    
    # Custom Styles
    title_style = ParagraphStyle(
        'TitleStyle',
        parent=styles['Heading1'],
        fontSize=22,
        textColor=colors.HexColor("#1e3a8a"),
        spaceAfter=6,
        alignment=1 # Center
    )
    subtitle_style = ParagraphStyle(
        'SubtitleStyle',
        parent=styles['Heading2'],
        fontSize=15,
        textColor=colors.HexColor("#3b82f6"),
        spaceAfter=15,
        alignment=1 # Center
    )
    meta_style = ParagraphStyle(
        'MetaStyle',
        parent=styles['Normal'],
        fontSize=9,
        textColor=colors.HexColor("#64748b")
    )
    
    elements = []
    
    # Branding Header
    elements.append(Paragraph("Yenova IT Club", title_style))
    
    query = db.query(Ledger)
    start_date, end_date, report_title = get_report_period(report_type)
    
    project = None
    if project_id:
        query = query.filter(Ledger.project_id == project_id)
        project = db.query(Project).filter(Project.id == project_id).first()
        if project:
            report_title = f"Project Financial Statement: {project.name}"
            safe_name = "".join(c if c.isalnum() or c in ('-', '_') else '_' for c in project.name)
            filename = f"{safe_name}_report_{datetime.now().strftime('%Y%m%d')}.pdf"
        else:
            filename = f"project_{project_id[:8]}_report_{datetime.now().strftime('%Y%m%d')}.pdf"
    else:
        if start_date and end_date:
            query = query.filter(
                Ledger.transaction_date >= start_date.date(),
                Ledger.transaction_date <= end_date.date()
            )
        filename = f"yenova_{report_type}_report_{datetime.now().strftime('%Y%m%d')}.pdf"
        
    elements.append(Paragraph(report_title, subtitle_style))
    
    # Date and User Info
    period_text = (
        f"Period: {start_date.strftime('%d %b %Y')} to {end_date.strftime('%d %b %Y')}"
        if (start_date and end_date)
        else f"Generated: {datetime.now().strftime('%d %B %Y %H:%M')}"
    )
    if project:
        period_text += f" | Code: {project.project_code} | Budget: Rs. {float(project.allocated_budget or 0):,.2f}"
        
    elements.append(Paragraph(period_text, meta_style))
    elements.append(Paragraph(f"Auditor / Generated By: {current_user.name} ({current_user.email})", meta_style))
    elements.append(Spacer(1, 15))
    
    transactions = query.order_by(Ledger.transaction_date.desc()).all()
    
    # Summary Metrics
    income = sum(t.amount for t in transactions if t.type == TransactionType.Income and t.status != TransactionStatus.VOIDED)
    expense = sum(t.amount for t in transactions if t.type == TransactionType.Expense and t.status != TransactionStatus.VOIDED)
    net_surplus = income - expense
    
    summary_data = [
        ['Metric', 'Amount (INR)'],
        ['Total Income / Inflow', f'Rs. {income:,.2f}'],
        ['Total Expenses / Outflow', f'Rs. {expense:,.2f}'],
        ['Net Balance / Surplus', f'Rs. {net_surplus:,.2f}'],
        ['Total Transactions', str(len(transactions))]
    ]
    if project:
        remaining_budget = float(project.allocated_budget or 0) - expense
        summary_data.insert(3, ['Allocated Budget', f'Rs. {float(project.allocated_budget or 0):,.2f}'])
        summary_data.insert(4, ['Remaining Budget', f'Rs. {remaining_budget:,.2f}'])
        
    summary_table = Table(summary_data, colWidths=[240, 180])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1e3a8a")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('TOPPADDING', (0, 0), (-1, 0), 6),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor("#f8fafc")),
        ('TEXTCOLOR', (0, 1), (-1, -1), colors.HexColor("#1e293b")),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 5),
        ('TOPPADDING', (0, 1), (-1, -1), 5)
    ]))
    
    elements.append(summary_table)
    elements.append(Spacer(1, 15))
    
    # Transaction Table
    elements.append(Paragraph(f"Transaction Ledger ({len(transactions)} Records)", styles['Heading3']))
    elements.append(Spacer(1, 5))
    
    table_data = [['Date', 'Txn #', 'Type', 'Description', 'Method', 'Amount (INR)']]
    for t in transactions[:45]:  # Up to 45 rows for PDF readability
        table_data.append([
            t.transaction_date.strftime('%Y-%m-%d') if t.transaction_date else "",
            t.transaction_number or "",
            t.type.value if hasattr(t.type, 'value') else str(t.type),
            (t.description[:26] + '..') if len(t.description) > 26 else t.description,
            t.payment_method or "-",
            f"{t.amount:,.2f}"
        ])
        
    t_table = Table(table_data, colWidths=[70, 75, 55, 175, 65, 80])
    t_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#3b82f6")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('ALIGN', (-1, 0), (-1, -1), 'RIGHT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor("#ffffff")),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0"))
    ]))
    
    elements.append(t_table)
    
    if len(transactions) > 45:
        elements.append(Spacer(1, 8))
        elements.append(Paragraph(f"Note: Displaying first 45 of {len(transactions)} transactions. For full dataset, please export to Excel (.xlsx) or CSV.", meta_style))
    
    # Build PDF
    doc.build(elements)
    
    buffer.seek(0)
    return StreamingResponse(
        iter([buffer.getvalue()]), 
        media_type="application/pdf", 
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"; filename*=UTF-8\'\'{filename}',
            "Access-Control-Expose-Headers": "Content-Disposition"
        }
    )

@router.get("/excel")
def export_excel(
    report_type: str = "all",
    project_id: str = None,
    db: Session = Depends(get_db), 
    current_user: User = Depends(get_current_user)
):
    wb = openpyxl.Workbook()
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    sub_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    bold_font = Font(bold=True)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    start_date, end_date, report_title = get_report_period(report_type)
    
    if project_id:
        project = db.query(Project).filter(Project.id == project_id).first()
        if not project:
            raise HTTPException(status_code=404, detail="Project not found")
            
        # Sheet 1: Project Financial Summary
        ws_summary = wb.active
        ws_summary.title = "Project Summary"
        ws_summary.append(["YENOVA IT CLUB - EVENT / PROJECT FINANCIAL REPORT", ""])
        ws_summary.append([])
        ws_summary.append(["Project Name", project.name])
        ws_summary.append(["Project Code", project.project_code])
        ws_summary.append(["Status", project.status.value if hasattr(project.status, 'value') else str(project.status)])
        ws_summary.append(["Allocated Budget (INR)", float(project.allocated_budget or 0)])
        
        # Calculate financial metrics
        transactions = db.query(Ledger).filter(Ledger.project_id == project_id).order_by(Ledger.transaction_date.desc()).all()
        income = sum(t.amount for t in transactions if t.type == TransactionType.Income and t.status != TransactionStatus.VOIDED)
        expense = sum(t.amount for t in transactions if t.type == TransactionType.Expense and t.status != TransactionStatus.VOIDED)
        remaining = float(project.allocated_budget or 0) - expense
        
        ws_summary.append(["Total Collected / Inflow (INR)", float(income)])
        ws_summary.append(["Total Spent / Outflow (INR)", float(expense)])
        ws_summary.append(["Remaining Budget (INR)", float(remaining)])
        ws_summary.append(["Net Profit / Surplus (INR)", float(income - expense)])
        ws_summary.append(["Transaction Count", len(transactions)])
        ws_summary.append([])
        ws_summary.append(["Generated By", current_user.name])
        ws_summary.append(["Export Date", datetime.now().strftime('%Y-%m-%d %H:%M:%S')])
        
        # Style Summary
        ws_summary.column_dimensions['A'].width = 32
        ws_summary.column_dimensions['B'].width = 30
        
        # Sheet 2: Detailed Transactions
        ws_tx = wb.create_sheet(title="Transactions")
        tx_headers = [
            "Date", "Transaction Number", "Type", "Status", "Description", 
            "Reference / UTR", "Payment Method", "Amount (INR)", "Custom Metadata"
        ]
        ws_tx.append(tx_headers)
        for col_num in range(1, len(tx_headers) + 1):
            cell = ws_tx.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin_border
            
        for t in transactions:
            ws_tx.append([
                t.transaction_date.strftime('%Y-%m-%d') if t.transaction_date else "",
                t.transaction_number,
                t.type.value if hasattr(t.type, 'value') else str(t.type),
                t.status.value if hasattr(t.status, 'value') else str(t.status),
                t.description,
                t.reference_number or "",
                t.payment_method or "",
                float(t.amount),
                t.custom_metadata or ""
            ])
            
        for col in ws_tx.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws_tx.column_dimensions[col_letter].width = max(max_len + 3, 12)
            
        safe_name = "".join(c if c.isalnum() or c in ('-', '_') else '_' for c in project.name)
        filename = f"{safe_name}_Report_{datetime.now().strftime('%Y%m%d')}.xlsx"
    else:
        query = db.query(Ledger)
        if start_date and end_date:
            query = query.filter(
                Ledger.transaction_date >= start_date.date(),
                Ledger.transaction_date <= end_date.date()
            )
        transactions = query.order_by(Ledger.transaction_date.desc()).all()
        
        income = sum(t.amount for t in transactions if t.type == TransactionType.Income and t.status != TransactionStatus.VOIDED)
        expense = sum(t.amount for t in transactions if t.type == TransactionType.Expense and t.status != TransactionStatus.VOIDED)
        net_surplus = income - expense
        
        # Sheet 1: Executive Summary
        ws_summary = wb.active
        ws_summary.title = "Executive Summary"
        ws_summary.append(["YENOVA IT CLUB - FINANCIAL STATEMENT", ""])
        ws_summary.append(["Report Scope", report_title])
        ws_summary.append([
            "Period Covered", 
            f"{start_date.strftime('%Y-%m-%d') if start_date else 'Beginning'} to {end_date.strftime('%Y-%m-%d') if end_date else datetime.now().strftime('%Y-%m-%d')}"
        ])
        ws_summary.append([])
        ws_summary.append(["Total Income / Inflow (INR)", float(income)])
        ws_summary.append(["Total Expense / Outflow (INR)", float(expense)])
        ws_summary.append(["Net Balance / Surplus (INR)", float(net_surplus)])
        ws_summary.append(["Transaction Count", len(transactions)])
        ws_summary.append([])
        ws_summary.append(["Generated By", current_user.name])
        ws_summary.append(["Generated At", datetime.now().strftime('%Y-%m-%d %H:%M:%S')])
        
        ws_summary.column_dimensions['A'].width = 30
        ws_summary.column_dimensions['B'].width = 35
        
        # Sheet 2: Ledger
        ws_tx = wb.create_sheet(title="Ledger Transactions")
        tx_headers = [
            "Date", "Transaction Number", "Type", "Status", "Description", 
            "Reference / UTR", "Payment Method", "Amount (INR)", "Custom Metadata"
        ]
        ws_tx.append(tx_headers)
        for col_num in range(1, len(tx_headers) + 1):
            cell = ws_tx.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin_border
            
        for t in transactions:
            ws_tx.append([
                t.transaction_date.strftime('%Y-%m-%d') if t.transaction_date else "",
                t.transaction_number,
                t.type.value if hasattr(t.type, 'value') else str(t.type),
                t.status.value if hasattr(t.status, 'value') else str(t.status),
                t.description,
                t.reference_number or "",
                t.payment_method or "",
                float(t.amount),
                t.custom_metadata or ""
            ])
            
        for col in ws_tx.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws_tx.column_dimensions[col_letter].width = max(max_len + 3, 12)
            
        filename = f"yenova_{report_type}_report_{datetime.now().strftime('%Y%m%d')}.xlsx"
        
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    return StreamingResponse(
        iter([buffer.getvalue()]), 
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", 
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"; filename*=UTF-8\'\'{filename}',
            "Access-Control-Expose-Headers": "Content-Disposition"
        }
    )

@router.get("/csv")
def export_csv(
    report_type: str = "all",
    project_id: str = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """Export transactions to standard CSV format."""
    query = db.query(Ledger)
    start_date, end_date, _ = get_report_period(report_type)
    
    if project_id:
        query = query.filter(Ledger.project_id == project_id)
        project = db.query(Project).filter(Project.id == project_id).first()
        safe_name = "".join(c if c.isalnum() or c in ('-', '_') else '_' for c in (project.name if project else "project"))
        filename = f"{safe_name}_data_{datetime.now().strftime('%Y%m%d')}.csv"
    else:
        if start_date and end_date:
            query = query.filter(
                Ledger.transaction_date >= start_date.date(),
                Ledger.transaction_date <= end_date.date()
            )
        filename = f"yenova_{report_type}_data_{datetime.now().strftime('%Y%m%d')}.csv"
        
    transactions = query.order_by(Ledger.transaction_date.desc()).all()
    
    output = io.StringIO()
    writer = csv.writer(output)
    
    # Headers
    writer.writerow([
        "Date",
        "Transaction Number",
        "Type",
        "Status",
        "Description",
        "Reference Number",
        "Payment Method",
        "Amount",
        "Custom Metadata"
    ])
    
    for t in transactions:
        writer.writerow([
            t.transaction_date.strftime('%Y-%m-%d') if t.transaction_date else "",
            t.transaction_number or "",
            t.type.value if hasattr(t.type, 'value') else str(t.type),
            t.status.value if hasattr(t.status, 'value') else str(t.status),
            t.description or "",
            t.reference_number or "",
            t.payment_method or "",
            t.amount,
            t.custom_metadata or ""
        ])
        
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue().encode('utf-8')]),
        media_type="text/csv",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"; filename*=UTF-8\'\'{filename}',
            "Access-Control-Expose-Headers": "Content-Disposition"
        }
    )

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

@router.get("/docx")
def export_docx(
    report_type: str = "yearly",
    project_id: str = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """Generate professional Word Document (.docx) financial reports."""
    doc = docx.Document()
    
    # Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    start_date, end_date, report_title = get_report_period(report_type)
    
    project = None
    if project_id:
        query = db.query(Ledger).filter(Ledger.project_id == project_id)
        project = db.query(Project).filter(Project.id == project_id).first()
        if project:
            report_title = f"Project Financial Report: {project.name}"
            safe_name = "".join(c if c.isalnum() or c in ('-', '_') else '_' for c in project.name)
            filename = f"{safe_name}_report_{datetime.now().strftime('%Y%m%d')}.docx"
        else:
            filename = f"project_{project_id[:8]}_report_{datetime.now().strftime('%Y%m%d')}.docx"
    else:
        query = db.query(Ledger)
        if start_date and end_date:
            query = query.filter(
                Ledger.transaction_date >= start_date.date(),
                Ledger.transaction_date <= end_date.date()
            )
        filename = f"yenova_{report_type}_report_{datetime.now().strftime('%Y%m%d')}.docx"

    transactions = query.order_by(Ledger.transaction_date.desc()).all()
    income = sum(t.amount for t in transactions if t.type == TransactionType.Income and t.status != TransactionStatus.VOIDED)
    expense = sum(t.amount for t in transactions if t.type == TransactionType.Expense and t.status != TransactionStatus.VOIDED)
    net_surplus = income - expense

    # Document Header
    p_title = doc.add_heading("Yenova IT Club", level=0)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p_title.runs:
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    p_sub = doc.add_heading(report_title, level=1)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p_sub.runs:
        run.font.color.rgb = RGBColor(0x3B, 0x82, 0xF6)

    # Metadata
    period_str = f"{start_date.strftime('%d %B %Y')} to {end_date.strftime('%d %B %Y')}" if (start_date and end_date) else f"Complete History up to {datetime.now().strftime('%d %B %Y')}"
    p_meta = doc.add_paragraph()
    p_meta.add_run(f"Period: {period_str}\n").bold = True
    p_meta.add_run(f"Auditor / Generated By: {current_user.name} ({current_user.email})\n")
    p_meta.add_run(f"Export Date: {datetime.now().strftime('%d %B %Y %H:%M:%S')}")
    if project:
        p_meta.add_run(f"\nProject Code: {project.project_code} | Allocated Budget: Rs. {float(project.allocated_budget or 0):,.2f}")

    # Summary KPI Table
    doc.add_heading("Executive Financial Summary", level=2)
    kpi_table = doc.add_table(rows=1, cols=2)
    kpi_table.style = 'Table Grid'
    hdr_cells = kpi_table.rows[0].cells
    hdr_cells[0].text = "Financial Metric"
    hdr_cells[1].text = "Amount (INR)"
    
    kpis = [
        ("Total Income / Inflow", f"Rs. {income:,.2f}"),
        ("Total Expense / Outflow", f"Rs. {expense:,.2f}"),
        ("Net Operating Surplus / Balance", f"Rs. {net_surplus:,.2f}"),
        ("Total Transaction Records", str(len(transactions)))
    ]
    if project:
        remaining_budget = float(project.allocated_budget or 0) - expense
        kpis.insert(2, ("Allocated Project Budget", f"Rs. {float(project.allocated_budget or 0):,.2f}"))
        kpis.insert(3, ("Remaining Project Budget", f"Rs. {remaining_budget:,.2f}"))

    for metric, val in kpis:
        row = kpi_table.add_row().cells
        row[0].text = metric
        row[1].text = val

    doc.add_paragraph() # Spacer

    # Transactions Table
    doc.add_heading("Itemized Transactions", level=2)
    tx_table = doc.add_table(rows=1, cols=6)
    tx_table.style = 'Table Grid'
    tx_hdrs = tx_table.rows[0].cells
    col_names = ["Date", "Txn #", "Type", "Description", "Method", "Amount (INR)"]
    for i, name in enumerate(col_names):
        tx_hdrs[i].text = name
        if tx_hdrs[i].paragraphs and tx_hdrs[i].paragraphs[0].runs:
            tx_hdrs[i].paragraphs[0].runs[0].font.bold = True

    for t in transactions[:60]: # up to 60 rows for Word document layout
        row = tx_table.add_row().cells
        row[0].text = t.transaction_date.strftime('%Y-%m-%d') if t.transaction_date else ""
        row[1].text = t.transaction_number or ""
        row[2].text = t.type.value if hasattr(t.type, 'value') else str(t.type)
        row[3].text = t.description or ""
        row[4].text = t.payment_method or "-"
        row[5].text = f"Rs. {t.amount:,.2f}"

    if len(transactions) > 60:
        p_note = doc.add_paragraph(f"Note: Displaying first 60 of {len(transactions)} records. For complete raw ledger records, please export via Excel (.xlsx) or CSV.")
        p_note.italic = True

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        iter([buffer.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"; filename*=UTF-8\'\'{filename}',
            "Access-Control-Expose-Headers": "Content-Disposition"
        }
    )

