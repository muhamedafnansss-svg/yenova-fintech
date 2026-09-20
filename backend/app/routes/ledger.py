import os
import io
import csv
import re
import json
from datetime import datetime, date
from fastapi import APIRouter, Depends, Query, HTTPException, UploadFile, File
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import or_, func
from typing import List, Optional
import openpyxl
import pdfplumber
import pypdf
from app.database.config import get_db
from app.models.ledger import Ledger, TransactionStatus, TransactionType
from app.models.category import Category
from app.models.project import Project
from app.schemas.ledger import LedgerResponse
from app.models.user import User
from app.models.role import Role
from app.middleware.deps import get_current_user, require_permission

router = APIRouter(prefix="/api/ledger", tags=["ledger"])

@router.get("/summary")
def get_ledger_summary(
    search: str = Query(None, description="Filter stats by search query"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    try:
        from app.services.balance import get_balance_details
        query = db.query(Ledger)
        if search:
            query = query.filter(
                or_(
                    Ledger.transaction_number.ilike(f"%{search}%"),
                    Ledger.description.ilike(f"%{search}%"),
                    Ledger.reference_number.ilike(f"%{search}%"),
                    Ledger.custom_metadata.ilike(f"%{search}%")
                )
            )

        total_records = query.count()

        # Inflow (all active income, excluding voided)
        total_income = float(
            query.filter(
                Ledger.type == TransactionType.Income,
                Ledger.status != TransactionStatus.VOIDED
            ).with_entities(func.sum(Ledger.amount)).scalar() or 0.0
        )

        # Outflow (all active expenses, excluding voided)
        total_expense = float(
            query.filter(
                Ledger.type == TransactionType.Expense,
                Ledger.status != TransactionStatus.VOIDED
            ).with_entities(func.sum(Ledger.amount)).scalar() or 0.0
        )

        verified_count = query.filter(Ledger.status == TransactionStatus.VERIFIED).count()
        voided_count = query.filter(Ledger.status == TransactionStatus.VOIDED).count()

        # Balance details (base opening balance + total net)
        bal_details = get_balance_details(db)
        opening_balance = bal_details.get("opening_balance", 0.0)
        live_balance = opening_balance + total_income - total_expense

        return {
            "total_records": total_records,
            "total_income": round(total_income, 2),
            "total_expense": round(total_expense, 2),
            "verified_count": verified_count,
            "voided_count": voided_count,
            "opening_balance": round(opening_balance, 2),
            "live_balance": round(live_balance, 2),
            "financial_year": bal_details.get("financial_year", "2026-2027")
        }
    except Exception as e:
        import traceback
        raise HTTPException(status_code=500, detail=traceback.format_exc())

@router.get("", response_model=List[LedgerResponse])
def get_ledger(
    skip: int = 0,
    limit: int = 100,
    search: str = Query(None, description="Search by transaction number, description, reference"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    try:
        query = db.query(Ledger).options(joinedload(Ledger.category), joinedload(Ledger.project))
        
        if search:
            query = query.filter(
                or_(
                    Ledger.transaction_number.ilike(f"%{search}%"),
                    Ledger.description.ilike(f"%{search}%"),
                    Ledger.reference_number.ilike(f"%{search}%"),
                    Ledger.custom_metadata.ilike(f"%{search}%")
                )
            )
            
        return query.order_by(Ledger.transaction_date.desc(), Ledger.created_at.desc()).offset(skip).limit(limit).all()
    except Exception as e:
        import traceback
        raise HTTPException(status_code=500, detail=traceback.format_exc())

@router.get("/{id}", response_model=LedgerResponse)
def get_ledger_entry(
    id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    entry = db.query(Ledger).filter(Ledger.id == id).first()
    if not entry:
        raise HTTPException(status_code=404, detail="Entry not found")
    return entry

from pydantic import BaseModel
from app.models.document import Document

class BulkDeleteRequest(BaseModel):
    transaction_ids: List[str]
    permanent: bool = True
    reason: Optional[str] = None

class BulkVerifyRequest(BaseModel):
    transaction_ids: List[str]

@router.post("/bulk-delete")
def bulk_delete_transactions(
    data: BulkDeleteRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    if not data.transaction_ids:
        raise HTTPException(status_code=400, detail="No transactions selected")
    
    # Check permissions
    role = db.query(Role).filter(Role.id == current_user.role_id).first()
    is_admin = (role and role.name == "Admin") or (current_user.role_id == 1)
    perms = (role.permissions if (role and role.permissions) else [])
    has_perm = is_admin or ("DELETE_EXPENSE" in perms or "VOID_TRANSACTION" in perms)
    if not has_perm:
        raise HTTPException(status_code=403, detail="Not authorized to delete transactions")
    
    deleted_ids = []
    from app.services.audit import AuditService

    for tid in data.transaction_ids:
        entry = db.query(Ledger).filter(Ledger.id == tid).first()
        if not entry:
            continue
            
        if data.permanent:
            # Unlink any documents before deleting
            db.query(Document).filter(Document.transaction_id == tid).update({"transaction_id": None})
            
            AuditService.log_action(
                db=db,
                user_id=current_user.uuid,
                action="DELETE_TRANSACTION",
                module="Ledger",
                entity_type="Transaction",
                entity_id=entry.id,
                old_values={"transaction_number": entry.transaction_number, "amount": entry.amount, "type": entry.type.value},
                new_values={"deleted": True, "reason": data.reason or "Bulk delete"}
            )
            db.delete(entry)
            deleted_ids.append(tid)
        else:
            # Void
            if entry.status != TransactionStatus.VOIDED:
                old_status = entry.status
                entry.status = TransactionStatus.VOIDED
                AuditService.log_action(
                    db=db,
                    user_id=current_user.uuid,
                    action="VOID_TRANSACTION",
                    module="Ledger",
                    entity_type="Transaction",
                    entity_id=entry.id,
                    old_values={"status": old_status.value},
                    new_values={"status": TransactionStatus.VOIDED.value, "reason": data.reason or "Bulk void"}
                )
                deleted_ids.append(tid)

    db.commit()
    action_str = "permanently deleted" if data.permanent else "voided"
    return {
        "message": f"{len(deleted_ids)} transaction(s) {action_str} successfully",
        "deleted_count": len(deleted_ids),
        "ids": deleted_ids
    }

@router.post("/bulk-verify")
def bulk_verify_transactions(
    data: BulkVerifyRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("VERIFY_TRANSACTION"))
):
    if not data.transaction_ids:
        raise HTTPException(status_code=400, detail="No transactions selected")
        
    verified_ids = []
    from app.services.audit import AuditService

    for tid in data.transaction_ids:
        entry = db.query(Ledger).filter(Ledger.id == tid).first()
        if entry and entry.status != TransactionStatus.VOIDED and entry.status != TransactionStatus.VERIFIED:
            old_status = entry.status
            entry.status = TransactionStatus.VERIFIED
            AuditService.log_action(
                db=db,
                user_id=current_user.uuid,
                action="VERIFY_TRANSACTION",
                module="Ledger",
                entity_type="Transaction",
                entity_id=entry.id,
                old_values={"status": old_status.value},
                new_values={"status": TransactionStatus.VERIFIED.value}
            )
            verified_ids.append(tid)

    db.commit()
    return {
        "message": f"{len(verified_ids)} transaction(s) verified successfully",
        "verified_count": len(verified_ids),
        "ids": verified_ids
    }

@router.delete("/{id}")
def delete_or_void_ledger_entry(
    id: str,
    permanent: bool = Query(True, description="If true, permanently deletes the transaction from database; otherwise voids it"),
    reason: Optional[str] = Query(None, description="Reason for voiding or deleting the transaction"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    entry = db.query(Ledger).filter(Ledger.id == id).first()
    if not entry:
        raise HTTPException(status_code=404, detail="Entry not found")
    
    role = db.query(Role).filter(Role.id == current_user.role_id).first()
    is_admin = (role and role.name == "Admin") or (current_user.role_id == 1)
    perms = (role.permissions if (role and role.permissions) else [])
    
    from app.services.audit import AuditService
    
    if permanent:
        if not is_admin and "DELETE_EXPENSE" not in perms:
            raise HTTPException(status_code=403, detail="Not authorized to delete transactions permanently")
            
        # Unlink any documents before deleting
        db.query(Document).filter(Document.transaction_id == id).update({"transaction_id": None})
        
        AuditService.log_action(
            db=db,
            user_id=current_user.uuid,
            action="DELETE_TRANSACTION",
            module="Ledger",
            entity_type="Transaction",
            entity_id=entry.id,
            old_values={"transaction_number": entry.transaction_number, "amount": entry.amount, "type": entry.type.value},
            new_values={"deleted": True, "reason": reason or "Deleted by user"}
        )
        db.delete(entry)
        db.commit()
        return {"message": "Transaction permanently deleted successfully", "deleted": True, "id": id}
    else:
        # Void logic
        if not is_admin and "VOID_TRANSACTION" not in perms:
            raise HTTPException(status_code=403, detail="Not authorized to void transactions")
            
        if entry.status == TransactionStatus.VOIDED:
            return {"message": "Entry is already voided", "status": "Voided", "id": id}
            
        old_status = entry.status
        entry.status = TransactionStatus.VOIDED
        
        AuditService.log_action(
            db=db,
            user_id=current_user.uuid,
            action="VOID_TRANSACTION",
            module="Ledger",
            entity_type="Transaction",
            entity_id=entry.id,
            old_values={"status": old_status.value},
            new_values={"status": TransactionStatus.VOIDED.value, "reason": reason or "Voided by user"}
        )
        db.commit()
        return {"message": "Ledger entry voided successfully", "status": "Voided", "id": id}

@router.post("/{id}/verify")
def verify_ledger_entry(
    id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("VERIFY_TRANSACTION"))
):
    entry = db.query(Ledger).filter(Ledger.id == id).first()
    if not entry:
        raise HTTPException(status_code=404, detail="Entry not found")
        
    if entry.status == TransactionStatus.VERIFIED:
        raise HTTPException(status_code=400, detail="Entry is already verified")
        
    old_status = entry.status
    entry.status = TransactionStatus.VERIFIED
    
    from app.services.audit import AuditService
    AuditService.log_action(
        db=db,
        user_id=current_user.uuid,
        action="VERIFY_TRANSACTION",
        module="Ledger",
        entity_type="Transaction",
        entity_id=entry.id,
        old_values={"status": old_status.value},
        new_values={"status": TransactionStatus.VERIFIED.value}
    )
    
    db.commit()
    return {"message": "Ledger entry verified successfully"}

import io
import re
import csv
import json
import uuid
from datetime import datetime, date
import openpyxl
from fastapi import UploadFile, File
from app.models.ledger import TransactionType
from app.services.audit import AuditService

def clean_amount(val):
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    # If the value looks like a date or time, do not treat as amount
    if re.search(r'\d{1,4}[-/]\d{1,2}[-/]\d{1,4}', s) or re.search(r'\d{1,2}:\d{2}', s):
        return 0.0
    s = re.sub(r'^[^\d.-]+', '', s)
    match = re.search(r'[\d,]+(?:\.\d+)?', s)
    if match:
        num_str = match.group(0).replace(',', '')
        try:
            return float(num_str)
        except ValueError:
            return 0.0
    return 0.0

GENERIC_REF_PLACEHOLDERS = {
    "", "none", "na", "n/a", "nil", "null", "no", "yes", "cash", "gpay", "google pay", 
    "phonepe", "paytm", "upi", "done", "paid", "ok", "offline", "hand", "direct", 
    "-", "--", "---", ".", "..", "completed", "success", "successful", "true", "false"
}

def is_placeholder_reference(ref_val: Optional[str]) -> bool:
    if not ref_val:
        return True
    cleaned = str(ref_val).strip().lower()
    if cleaned in GENERIC_REF_PLACEHOLDERS:
        return True
    if len(cleaned) < 4:
        return True
    return False

def detect_columns(headers, sample_rows):
    lower_headers = {h: h.strip().lower() for h in headers}
    
    # 1. Amount
    amount_keywords = [
        'amount', 'paid', 'fee', 'price', 'total', 'cost', 'rs', 'inr', 'rupees', 
        'amt', 'charge', 'reg fee', 'registration fee', 'ticket', 'debit', 'credit',
        'withdrawal', 'deposit', 'sum', 'value', 'payment amount'
    ]
    negative_amount_keywords = [
        'timestamp', 'time', 'date', 'datetime', 'created', 'updated', 'hall', 'lh', 
        'roll', 'reg no', 'reg_no', 'phone', 'mobile', 'contact', 'room', 'sem', 
        'year', 'id', 'ref', 'pin', 'code', 'zip', 'no', 'num', 'screenshot', 
        'proof', 'photo', 'count', 'sl', 'sl no', 'serial'
    ]
    
    amount_col = None
    for h, lh in lower_headers.items():
        if any(kw in lh for kw in amount_keywords):
            if not any(neg in lh for neg in ['id', 'ref', 'token', 'hall', 'lh', 'phone', 'roll', 'time', 'date', 'timestamp']):
                amount_col = h
                break
            elif any(strong in lh for strong in ['fee', 'paid', 'price', 'amount', 'amt', 'total']):
                amount_col = h
                break

    # Only guess by numeric values if the column does NOT match any negative keywords
    if not amount_col and sample_rows:
        for h, lh in lower_headers.items():
            if any(neg in lh for neg in negative_amount_keywords):
                continue
            vals = [clean_amount(r.get(h)) for r in sample_rows[:10]]
            if any(v >= 10 for v in vals):
                amount_col = h
                break

    # 2. Payer / Description
    name_keywords = [
        'name', 'student', 'payer', 'participant', 'attendee', 'member', 'person', 
        'full_name', 'candidate', 'user', 'client', 'description', 'particulars', 
        'narration', 'purpose', 'details', 'item', 'vendor', 'payee'
    ]
    name_col = None
    for h, lh in lower_headers.items():
        if any(kw in lh for kw in name_keywords) and h != amount_col:
            name_col = h
            break

    # 3. Reference / UTR / Transaction ID
    strong_ref_keywords = [
        'utr', 'transaction_id', 'transaction id', 'txn_id', 'txn id', 
        'upi transaction id', 'upi id', 'upi_id', 'upi', 'scanner_ref', 
        'reference_number', 'reference number', 'ref_no', 'ref no', 'ref',
        'receipt_no', 'receipt no', 'order_id', 'order id', 'chq', 'cheque'
    ]
    general_ref_keywords = ['ref', 'txn', 'trans', 'reference', 'receipt', 'order_id', 'order id']
    
    ref_col = None
    for h, lh in lower_headers.items():
        if any(kw in lh for kw in strong_ref_keywords) and h not in [amount_col, name_col]:
            ref_col = h
            break

    if not ref_col:
        for h, lh in lower_headers.items():
            if any(kw in lh for kw in general_ref_keywords) and h not in [amount_col, name_col]:
                ref_col = h
                break

    if not ref_col:
        for h, lh in lower_headers.items():
            if 'id' in lh and not any(neg in lh for neg in ['student', 'campus', 'roll', 'reg', 'register', 'user', 'member']) and h not in [amount_col, name_col]:
                ref_col = h
                break

    # 4. Date
    date_keywords = ['date', 'time', 'timestamp', 'created', 'day', 'datetime', 'txn date', 'value date']
    date_col = None
    for h, lh in lower_headers.items():
        if any(kw in lh for kw in date_keywords) and h not in [amount_col, name_col, ref_col]:
            date_col = h
            break

    # 5. Method
    method_keywords = ['payment_method', 'payment method', 'payment_mode', 'payment mode', 'method', 'channel', 'mode', 'via']
    neg_method = ['whatsapp', 'group', 'join', 'status', 'apply', 'registration']
    method_col = None
    for h, lh in lower_headers.items():
        if any(kw in lh for kw in method_keywords) and not any(neg in lh for neg in neg_method) and h not in [amount_col, name_col, ref_col, date_col]:
            method_col = h
            break

    # 6. Category
    cat_keywords = ['category', 'category_name', 'cat', 'expense_category', 'income_category', 'fee_type', 'ticket_type']
    cat_col = None
    for h, lh in lower_headers.items():
        if any(kw in lh for kw in cat_keywords) and h not in [amount_col, name_col, ref_col, date_col, method_col]:
            cat_col = h
            break

    # 7. Project / Event
    proj_keywords = ['project', 'project_name', 'event', 'event_name', 'campaign']
    proj_col = None
    for h, lh in lower_headers.items():
        if any(kw in lh for kw in proj_keywords) and h not in [amount_col, name_col, ref_col, date_col, method_col, cat_col]:
            proj_col = h
            break

    known_cols = set(filter(None, [amount_col, name_col, ref_col, date_col, method_col, cat_col, proj_col]))
    custom_cols = [h for h in headers if h not in known_cols]

    return {
        "amount_column": amount_col,
        "name_column": name_col,
        "reference_column": ref_col,
        "date_column": date_col,
        "payment_method_column": method_col,
        "category_column": cat_col,
        "project_column": proj_col,
        "custom_columns": custom_cols
    }


def extract_pdf_data(contents: bytes, filename: str) -> List[List[str]]:
    raw_rows = []
    
    # Strategy 1: Explicit bordered tables with pdfplumber
    try:
        with pdfplumber.open(io.BytesIO(contents)) as pdf:
            for p_idx, page in enumerate(pdf.pages):
                tables = page.extract_tables({"vertical_strategy": "lines", "horizontal_strategy": "lines"}) or []
                if not tables:
                    tables = page.extract_tables() or []
                for t in tables:
                    cleaned_t = []
                    for row in t:
                        if row and any(c is not None and str(c).strip() for c in row):
                            cleaned_t.append([str(c).strip() if c is not None else "" for c in row])
                    if cleaned_t and len(cleaned_t[0]) >= 2:
                        for r_idx, r in enumerate(cleaned_t):
                            if p_idx > 0 and r_idx == 0 and raw_rows and r == raw_rows[0]:
                                continue
                            raw_rows.append(r)
    except Exception:
        pass

    # Strategy 2: If no bordered multi-column table was found, extract full text and analyze structure
    if not raw_rows or len(raw_rows) < 2 or max(len(r) for r in raw_rows) < 2:
        raw_rows = []
        full_text = ""
        try:
            with pdfplumber.open(io.BytesIO(contents)) as pdf:
                for page in pdf.pages:
                    txt = page.extract_text() or ""
                    if txt:
                        full_text += txt + "\n"
        except Exception:
            pass

        if not full_text.strip():
            try:
                reader = pypdf.PdfReader(io.BytesIO(contents))
                for page in reader.pages:
                    txt = page.extract_text() or ""
                    if txt:
                        full_text += txt + "\n"
            except Exception:
                pass

        lines = [line.strip() for line in full_text.split("\n") if line.strip()]

        # Strategy 2A: Check for Key-Value Receipt / Single Voucher / Bill / Invoice
        has_colon_pairs = sum(1 for l in lines if ":" in l) >= 2
        receipt_kw = any(w in full_text.lower() for w in ["receipt", "invoice", "voucher", "payment", "paid", "bill", "ticket", "fee", "collection"])
        amount_match = re.search(r'(?:total\s*amount|amount\s*paid|grand\s*total|net\s*amount|total|amount|paid|fee|price|rs\.?|inr|₹)[\s:_-]*([0-9,]+(?:\.[0-9]{1,2})?)', full_text, re.IGNORECASE)
        amount_val = "0.00"
        if amount_match:
            amount_val = amount_match.group(1).replace(",", "")

        if (has_colon_pairs or receipt_kw) and float(amount_val or 0) > 0:
            kv = {}
            for l in lines:
                if ":" in l:
                    k, v = l.split(":", 1)
                    kv[k.strip().lower()] = v.strip()
            
            payer_val = ""
            for k in ["payer name", "payer", "name", "student name", "student", "customer", "received from", "paid by", "bill to"]:
                if k in kv:
                    payer_val = kv[k]
                    break
            if not payer_val:
                name_match = re.search(r'(?:student\s*name|payer\s*name|customer\s*name|name|payer|student|received\s*from|paid\s*by)[\s:_-]*([A-Za-z\s.]{3,35})', full_text, re.IGNORECASE)
                if name_match:
                    payer_val = name_match.group(1).strip()
            if not payer_val:
                payer_val = f"Receipt from {filename}"

            ref_val = ""
            for k in ["transaction ref / utr", "transaction ref", "utr", "transaction id", "ref no", "receipt no", "order id", "reference number", "reference"]:
                if k in kv:
                    ref_val = kv[k]
                    break
            if not ref_val:
                ref_match = re.search(r'(?:utr|txn|ref|order|receipt)[\s:_-]*([A-Za-z0-9_-]{6,30})', full_text, re.IGNORECASE)
                if ref_match:
                    ref_val = ref_match.group(1).strip()
            if not ref_val:
                ref_val = f"PDF-{datetime.now().strftime('%m%d%H%M%S')}"

            date_val = datetime.now().strftime("%Y-%m-%d")
            for k in ["date", "transaction date", "payment date"]:
                if k in kv:
                    date_val = kv[k]
                    break
            if not date_val or date_val == datetime.now().strftime("%Y-%m-%d"):
                date_match = re.search(r'([0-9]{1,4}[-/][0-9]{1,2}[-/][0-9]{1,4})', full_text)
                if date_match:
                    date_val = date_match.group(1)

            mode_val = "Scanner / UPI"
            for k in ["payment mode", "mode", "method", "payment method"]:
                if k in kv:
                    mode_val = kv[k]
                    break

            desc_val = kv.get("purpose") or kv.get("description") or f"Payment from {filename}"

            headers = ["Participant / Payer Name", "Payment Amount", "Transaction ID / UTR", "Transaction Date", "Payment Mode", "Purpose / Description"]
            row = [payer_val, amount_val, ref_val, date_val, mode_val, desc_val]
            raw_rows = [headers, row]

        # Strategy 2B: Multi-line Statement (Lines with Dates & Amounts, or Tabular Layout)
        if not raw_rows:
            stmt_rows = []
            stmt_headers = ["Transaction Date", "Description / Payer", "Payment Amount", "Reference Number", "Status"]
            for l in lines:
                d_match = re.search(r'([0-9]{1,4}[-/][0-9]{1,2}[-/][0-9]{1,4})', l)
                a_match = re.search(r'(?:\s|^)([0-9,]+\.[0-9]{2})(?:\s|$)', l)
                if d_match and a_match:
                    d_str = d_match.group(1)
                    amt_str = a_match.group(1).replace(",", "")
                    d_end = d_match.end()
                    a_start = a_match.start()
                    desc = l[d_end:a_start].strip() if a_start > d_end else "Transaction"
                    remainder = l[a_match.end():].strip()
                    ref = ""
                    status = "Completed"
                    parts = remainder.split()
                    if len(parts) >= 2:
                        ref = parts[0]
                        status = parts[1]
                    elif len(parts) == 1:
                        ref = parts[0]
                    stmt_rows.append([d_str, desc, amt_str, ref, status])

            if len(stmt_rows) >= 1:
                raw_rows = [stmt_headers] + stmt_rows

        # Strategy 2C: Delimited lines (pipes, tabs, commas)
        if not raw_rows:
            for delim in ["|", "\t", ","]:
                delim_rows = []
                for l in lines:
                    if delim in l:
                        parts = [p.strip() for p in l.split(delim)]
                        if len(parts) >= 2:
                            delim_rows.append(parts)
                if len(delim_rows) >= 2:
                    raw_rows = delim_rows
                    break

        # Strategy 2D: Scanned Receipt image inside PDF
        if not raw_rows:
            clean_fname = os.path.splitext(filename)[0]
            found_amounts = re.findall(r'(?:rs|inr|amt|fee)?[:\s_-]*([0-9]+(?:[.,][0-9]{2})?)', clean_fname + " " + full_text, re.IGNORECASE)
            amount_val = found_amounts[0] if found_amounts else "0.00"
            found_utr = re.findall(r'(?:utr|txn|ref|token)[:\s_-]*([A-Za-z0-9]{6,22})', clean_fname + " " + full_text, re.IGNORECASE)
            ref_val = found_utr[0] if found_utr else f"PDF-{datetime.now().strftime('%m%d%H%M')}"
            headers = ["Participant Name / Description", "Registration / Payment Amount", "Scanner Token / UTR", "Transaction Date", "Payment Mode", "Attached Receipt"]
            row_data = [f"Receipt from {filename}", amount_val, ref_val, datetime.now().strftime("%Y-%m-%d"), "Scanner / UPI", filename]
            raw_rows = [headers, row_data]

    return raw_rows


def extract_rows_from_bytes(contents: bytes, filename: str) -> List[List[str]]:
    raw_rows: List[List[str]] = []
    fname_lower = filename.lower()

    # 1. ZIP Archive (Modern .xlsx, Word .docx, or ZIP bundle)
    if contents.startswith(b"PK\x03\x04"):
        try:
            import zipfile
            zf = zipfile.ZipFile(io.BytesIO(contents))
            namelist = zf.namelist()

            # A. Check if it's an Excel .xlsx
            if any("xl/workbook" in n or "xl/worksheets" in n for n in namelist):
                try:
                    wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
                    for sheet in wb.worksheets:
                        s_rows = []
                        for row in sheet.iter_rows(values_only=True):
                            if row and any(c is not None and str(c).strip() for c in row):
                                s_rows.append([str(c).strip() if c is not None else "" for c in row])
                        if len(s_rows) > len(raw_rows):
                            raw_rows = s_rows
                except Exception:
                    pass

            # B. Check if it's a Word .docx
            elif any("word/document.xml" in n for n in namelist):
                try:
                    import docx
                    doc = docx.Document(io.BytesIO(contents))
                    for table in doc.tables:
                        for row in table.rows:
                            cells = [c.text.strip() for c in row.cells]
                            cleaned = []
                            for i, c in enumerate(cells):
                                if i == 0 or c != cells[i-1]:
                                    cleaned.append(c)
                            if any(cleaned):
                                raw_rows.append(cleaned)
                except Exception:
                    pass

            # C. Generic ZIP bundle: find any embedded spreadsheet/pdf/docx
            else:
                for name in namelist:
                    if name.lower().endswith((".xlsx", ".xls", ".csv", ".pdf", ".docx")):
                        inner_data = zf.read(name)
                        inner_rows = extract_rows_from_bytes(inner_data, name)
                        if len(inner_rows) > len(raw_rows):
                            raw_rows = inner_rows
        except Exception:
            pass

    # 2. PDF Documents (Magic bytes %PDF- or .pdf extension or fallback)
    if not raw_rows and (contents.startswith(b"%PDF-") or fname_lower.endswith(".pdf")):
        raw_rows = extract_pdf_data(contents, filename)

    # 3. Try openpyxl directly
    if not raw_rows:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
            for sheet in wb.worksheets:
                s_rows = []
                for row in sheet.iter_rows(values_only=True):
                    if row and any(c is not None and str(c).strip() for c in row):
                        s_rows.append([str(c).strip() if c is not None else "" for c in row])
                if len(s_rows) > len(raw_rows):
                    raw_rows = s_rows
        except Exception:
            pass

    # 4. Try Word document directly
    if not raw_rows:
        try:
            import docx
            doc = docx.Document(io.BytesIO(contents))
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    cleaned = []
                    for i, c in enumerate(cells):
                        if i == 0 or c != cells[i-1]:
                            cleaned.append(c)
                    if any(cleaned):
                        raw_rows.append(cleaned)
        except Exception:
            pass

    # 5. Try xlrd for legacy .xls / BIFF8
    if not raw_rows:
        try:
            import xlrd
            wb = xlrd.open_workbook(file_contents=contents)
            for s_idx in range(wb.nsheets):
                sheet = wb.sheet_by_index(s_idx)
                s_rows = []
                for r_idx in range(sheet.nrows):
                    row = sheet.row_values(r_idx)
                    if any(str(c).strip() for c in row):
                        s_rows.append([str(c).strip() for c in row])
                if len(s_rows) > len(raw_rows):
                    raw_rows = s_rows
        except Exception:
            pass

    # 6. Try HTML table parsing (Google Sheets web export, ERP exports)
    if not raw_rows:
        try:
            text = None
            for enc in ["utf-8-sig", "utf-8", "latin-1", "cp1252"]:
                try:
                    text = contents.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            if text and ("<table" in text.lower() or "<tr" in text.lower()):
                for tr in re.findall(r'<tr[^>]*>(.*?)</tr>', text, re.DOTALL | re.IGNORECASE):
                    cells = re.findall(r'<(?:td|th)[^>]*>(.*?)</(?:td|th)>', tr, re.DOTALL | re.IGNORECASE)
                    clean = [re.sub(r'<[^>]+>', '', c).strip() for c in cells]
                    if any(clean):
                        raw_rows.append(clean)
        except Exception:
            pass

    # 7. Try CSV / Delimited
    if not raw_rows:
        try:
            text = None
            for enc in ["utf-8-sig", "utf-8", "latin-1", "cp1252"]:
                try:
                    text = contents.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            if text is None:
                text = contents.decode("utf-8-sig", errors="ignore")
            sample = text[:2048]
            delim = ','
            for c in [',', '\t', ';', '|']:
                if sample.count(c) > sample.count(delim):
                    delim = c
            reader = csv.reader(io.StringIO(text), delimiter=delim)
            for row in reader:
                if row and any(str(c).strip() for c in row):
                    raw_rows.append([str(c).strip() for c in row])
        except Exception:
            pass

    # 8. Receipt Image Check
    if not raw_rows and any(fname_lower.endswith(img_ext) for img_ext in [".png", ".jpg", ".jpeg", ".webp"]):
        try:
            clean_fname = os.path.splitext(filename)[0]
            found_amounts = re.findall(r'(?:rs|inr|amt|fee)?[:\s_-]*([0-9]+(?:[.,][0-9]{2})?)', clean_fname, re.IGNORECASE)
            amount_val = found_amounts[0] if found_amounts else "0.00"
            found_utr = re.findall(r'(?:utr|txn|ref|token)[:\s_-]*([A-Za-z0-9]{6,22})', clean_fname, re.IGNORECASE)
            ref_val = found_utr[0] if found_utr else f"IMG-{datetime.now().strftime('%m%d%H%M')}"
            headers = ["Participant Name / Description", "Registration / Payment Amount", "Scanner Token / UTR", "Transaction Date", "Payment Mode", "Attached Receipt"]
            row_data = [f"Receipt proof from {filename}", amount_val, ref_val, datetime.now().strftime("%Y-%m-%d"), "Scanner / UPI", filename]
            raw_rows = [headers, row_data]
        except Exception:
            pass

    return raw_rows


def normalize_text_for_comparison(text: Optional[str]) -> str:
    if not text:
        return ""
    cleaned = re.sub(r'[^\w\s]', ' ', str(text).lower())
    return ' '.join(cleaned.split())


def extract_payer_from_desc(desc: str) -> str:
    if not desc:
        return ""
    if " | " in desc:
        return desc.split(" | ")[0].strip()
    lower_desc = desc.lower()
    for prefix in ["recieved from ", "received from ", "payment from ", "receipt from "]:
        if lower_desc.startswith(prefix):
            return desc[len(prefix):].strip()
    return desc.strip()


def build_duplicate_detector(db: Session) -> dict:
    """Build fast lookup indices across existing ledger entries for UTR, student IDs, phone, and participant names."""
    db_records = db.query(
        Ledger.id,
        Ledger.transaction_number,
        Ledger.reference_number,
        Ledger.amount,
        Ledger.description,
        Ledger.custom_metadata,
        Ledger.type
    ).filter(Ledger.status != TransactionStatus.VOIDED).all()

    existing_refs = {}
    existing_ids = {}
    existing_phones = {}
    existing_name_amounts = {}
    all_name_amounts = []

    for row in db_records:
        t_num = row.transaction_number or ""
        ref = row.reference_number
        if ref and not is_placeholder_reference(ref):
            existing_refs[str(ref).strip().lower()] = (t_num, str(ref).strip())

        desc = row.description or ""
        meta_str = row.custom_metadata or ""
        amt = float(row.amount) if row.amount is not None else 0.0

        payer = extract_payer_from_desc(desc)

        meta = {}
        if meta_str:
            try:
                meta = json.loads(meta_str)
            except Exception:
                pass

        # Extract student / campus / roll IDs from metadata
        for k, v in meta.items():
            k_lower = k.lower()
            v_str = str(v).strip()
            if not v_str:
                continue
            if any(id_kw in k_lower for id_kw in ['student', 'campus', 'roll', 'reg', 'register', 'admission', 'usn', 'prn']):
                clean_id = re.sub(r'[^A-Za-z0-9]', '', v_str).lower()
                if len(clean_id) >= 3 and (not clean_id.isdigit() or len(clean_id) >= 4):
                    existing_ids[clean_id] = (t_num, v_str, payer)
            if any(p_kw in k_lower for p_kw in ['phone', 'mobile', 'contact', 'whatsapp']):
                digits = re.sub(r'\D', '', v_str)
                if len(digits) >= 10:
                    existing_phones[digits[-10:]] = (t_num, payer)
            if any(n_kw in k_lower for n_kw in ['name', 'payer', 'participant']) and not any(neg in k_lower for neg in ['id', 'roll', 'hall', 'dept', 'department', 'screenshot']):
                if v_str:
                    payer = v_str

        # Extract IDs from description
        id_matches = re.findall(r'(?:campus\s*id|reg(?:istration)?\s*(?:no|number)?|roll\s*no)[:\s_-]*([A-Za-z0-9/_-]+)', desc, re.IGNORECASE)
        for im in id_matches:
            clean_id = re.sub(r'[^A-Za-z0-9]', '', im).lower()
            if len(clean_id) >= 3 and (not clean_id.isdigit() or len(clean_id) >= 4):
                existing_ids[clean_id] = (t_num, im, payer)

        # Phone numbers from description
        phone_matches = re.findall(r'(?:\+?91[\s-]?)?([6-9]\d{9})\b', desc)
        for pm in phone_matches:
            existing_phones[pm[-10:]] = (t_num, payer)

        # UTRs from description
        desc_refs = re.findall(r'(?:utr|txn|ref|token)[:\s_-]*([A-Za-z0-9]{8,22})', desc, re.IGNORECASE)
        for dr in desc_refs:
            if not is_placeholder_reference(dr):
                existing_refs[dr.lower()] = (t_num, dr)

        clean_p = normalize_text_for_comparison(payer)
        if clean_p and len(clean_p) >= 3 and amt > 0:
            existing_name_amounts[(clean_p, round(amt, 2))] = (t_num, payer)
            all_name_amounts.append((clean_p, round(amt, 2), t_num, payer))

    return {
        "existing_refs": existing_refs,
        "existing_ids": existing_ids,
        "existing_phones": existing_phones,
        "existing_name_amounts": existing_name_amounts,
        "all_name_amounts": all_name_amounts,
        "seen_refs_in_file": {},
        "seen_ids_in_file": {},
        "seen_phones_in_file": {},
        "seen_name_amounts_in_file": {}
    }


def detect_row_duplicate(
    row: dict,
    amount_col: Optional[str],
    name_col: Optional[str],
    ref_col: Optional[str],
    fixed_amt: Optional[float],
    detector: dict,
    row_idx: int
) -> tuple[bool, str]:
    """Evaluates multi-signal duplicate criteria: UTR, Student/Campus ID, Phone, and Name+Amount."""
    amt = 0.0
    if fixed_amt is not None and fixed_amt > 0:
        amt = fixed_amt
    elif amount_col and amount_col != "__FIXED__":
        amt = clean_amount(row.get(amount_col))
    if amt <= 0 and row.get("_parsed_amount"):
        amt = clean_amount(row.get("_parsed_amount"))

    # 1. Reference / UTR
    raw_ref = str(row.get(ref_col) or "").strip() if ref_col else ""
    if not raw_ref:
        for k, v in row.items():
            if any(kw in k.lower() for kw in ['utr', 'txn', 'transaction id', 'reference', 'upi transaction id']):
                if str(v).strip():
                    raw_ref = str(v).strip()
                    break

    if raw_ref and not is_placeholder_reference(raw_ref):
        ref_lower = raw_ref.lower()
        if ref_lower in detector["existing_refs"]:
            txn_num, orig_ref = detector["existing_refs"][ref_lower]
            return True, f"Duplicate Reference/UTR '{orig_ref}' (already in Ledger {txn_num})"
        if ref_lower in detector["seen_refs_in_file"]:
            prev_row = detector["seen_refs_in_file"][ref_lower]
            return True, f"Duplicate Reference/UTR in file (matches row #{prev_row})"
        detector["seen_refs_in_file"][ref_lower] = row_idx

    # 2. Student Registration / Campus ID / Roll Number
    for k, v in row.items():
        if k.startswith("_"):
            continue
        k_lower = k.lower()
        v_str = str(v).strip()
        if not v_str:
            continue
        if any(id_kw in k_lower for id_kw in ['student', 'campus', 'roll', 'reg', 'register', 'admission', 'usn', 'prn']):
            clean_id = re.sub(r'[^A-Za-z0-9]', '', v_str).lower()
            if len(clean_id) >= 3 and (not clean_id.isdigit() or len(clean_id) >= 4):
                if clean_id in detector["existing_ids"]:
                    txn_num, orig_id, orig_payer = detector["existing_ids"][clean_id]
                    return True, f"Duplicate Student/Campus ID '{orig_id}' (recorded in {txn_num} - {orig_payer})"
                if clean_id in detector["seen_ids_in_file"]:
                    prev_row = detector["seen_ids_in_file"][clean_id]
                    return True, f"Duplicate Student ID '{v_str}' in file (matches row #{prev_row})"
                detector["seen_ids_in_file"][clean_id] = row_idx

    # 3. Phone Number
    for k, v in row.items():
        if k.startswith("_"):
            continue
        v_str = str(v).strip()
        if not v_str:
            continue
        digits = re.sub(r'\D', '', v_str)
        if len(digits) >= 10:
            last10 = digits[-10:]
            if last10 in detector["existing_phones"]:
                txn_num, orig_payer = detector["existing_phones"][last10]
                return True, f"Duplicate Phone '{last10}' (already in {txn_num} - {orig_payer})"
            if last10 in detector["seen_phones_in_file"]:
                prev_row = detector["seen_phones_in_file"][last10]
                return True, f"Duplicate Phone in file (matches row #{prev_row})"
            detector["seen_phones_in_file"][last10] = row_idx

    # 4. Participant Name + Amount
    raw_name = str(row.get(name_col) or "").strip() if name_col else ""
    if not raw_name:
        for k, v in row.items():
            if any(n_kw in k.lower() for n_kw in ['full name', 'student name', 'name', 'participant', 'payer']) and not any(neg in k.lower() for neg in ['id', 'roll', 'hall', 'dept', 'department', 'screenshot']):
                if str(v).strip():
                    raw_name = str(v).strip()
                    break

    clean_name = normalize_text_for_comparison(raw_name)
    if clean_name and len(clean_name) >= 3 and amt > 0:
        key = (clean_name, round(amt, 2))
        if key in detector["existing_name_amounts"]:
            txn_num, orig_name = detector["existing_name_amounts"][key]
            return True, f"Duplicate Participant: '{orig_name}' already recorded with ₹{amt:,.2f} ({txn_num})"
        if key in detector["seen_name_amounts_in_file"]:
            prev_row = detector["seen_name_amounts_in_file"][key]
            return True, f"Duplicate Participant '{raw_name}' in file (matches row #{prev_row})"
        detector["seen_name_amounts_in_file"][key] = row_idx

        # 5. Fuzzy Name similarity with same amount
        words = set(clean_name.split())
        for ex_name, ex_amt, txn_num, orig_name in detector["all_name_amounts"]:
            if abs(ex_amt - round(amt, 2)) < 0.01:
                ex_words = set(ex_name.split())
                if len(words) >= 2 and len(ex_words) >= 2:
                    overlap = len(words & ex_words) / max(len(words), len(ex_words))
                    if overlap >= 0.70:
                        return True, f"Similar to '{orig_name}' with ₹{amt:,.2f} ({txn_num})"
                elif len(words) == 1 and len(ex_words) == 1:
                    if len(clean_name) >= 5 and (clean_name in ex_name or ex_name in clean_name):
                        return True, f"Similar to '{orig_name}' with ₹{amt:,.2f} ({txn_num})"

    return False, ""


@router.post("/import/parse")
async def parse_excel_payments(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    filename = file.filename or "uploaded_payment_file"
    contents = await file.read()
    if not contents:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")

    raw_rows = extract_rows_from_bytes(contents, filename)

    if not raw_rows:
        raise HTTPException(
            status_code=400, 
            detail=f"Could not extract payment records from '{filename}'. Please ensure the document contains tabular transaction data, key-value payment details, or a valid spreadsheet/PDF."
        )

    # Intelligent Header Detection
    header_index = 0
    max_cols = max(len([c for c in r if c]) for r in raw_rows[:10]) if raw_rows else 0
    for idx, r in enumerate(raw_rows[:10]):
        non_empty = len([c for c in r if c])
        if non_empty >= max(2, max_cols - 2):
            header_index = idx
            break

    headers = [str(c).strip() if str(c).strip() else f"Column_{i+1}" for i, c in enumerate(raw_rows[header_index])]

    rows = []
    for r in raw_rows[header_index + 1:]:
        row_dict = {}
        for i, h in enumerate(headers):
            val = r[i] if i < len(r) else ""
            row_dict[h] = val
        if any(row_dict.values()):
            rows.append(row_dict)

    if not rows and len(raw_rows) == 1:
        headers = [f"Field_{i+1}" for i in range(len(raw_rows[0]))]
        rows = [{headers[i]: raw_rows[0][i] for i in range(len(headers))}]

    if not headers or not rows:
        raise HTTPException(
            status_code=400, 
            detail=f"Detected headers {headers} but found no valid payment rows in '{filename}'."
        )

    mapping = detect_columns(headers, rows)
    amount_col = mapping.get("amount_column")
    name_col = mapping.get("name_column")
    ref_col = mapping.get("reference_column")

    # Build multi-signal duplicate detector
    detector = build_duplicate_detector(db)

    total_amount = 0.0
    duplicate_count = 0
    enhanced_rows = []

    for idx, r in enumerate(rows):
        amt = clean_amount(r.get(amount_col)) if amount_col else 0.0
        # If amount_col wasn't detected, check if any field has clean amount
        if amt <= 0:
            for k, v in r.items():
                parsed_val = clean_amount(v)
                if parsed_val > 0:
                    amt = parsed_val
                    break

        total_amount += amt
        
        is_dup, dup_reason = detect_row_duplicate(r, amount_col, name_col, ref_col, None, detector, idx + 1)
        if is_dup:
            duplicate_count += 1

        enhanced_rows.append({
            "_row_id": idx + 1,
            "_is_duplicate": is_dup,
            "_duplicate_reason": dup_reason,
            "_parsed_amount": amt,
            **r
        })

    return {
        "filename": filename,
        "headers": headers,
        "detected_mapping": mapping,
        "custom_columns": mapping["custom_columns"],
        "total_rows": len(enhanced_rows),
        "total_amount": round(total_amount, 2),
        "duplicate_count": duplicate_count,
        "rows": enhanced_rows
    }

class ImportCommitRequest(BaseModel):
    rows: List[dict]
    mapping: dict
    category_id: Optional[str] = None
    project_id: Optional[str] = None
    default_payment_method: Optional[str] = "Scanner / UPI"
    skip_duplicates: bool = True
    fixed_amount: Optional[float] = None
    amount_override: Optional[float] = None
    transaction_type: Optional[str] = "Income"

@router.post("/import/commit")
def commit_excel_payments(
    payload: ImportCommitRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    mapping = payload.mapping or {}
    amount_col = mapping.get("amount_column")
    name_col = mapping.get("name_column")
    ref_col = mapping.get("reference_column")
    date_col = mapping.get("date_column")
    method_col = mapping.get("payment_method_column")
    category_col = mapping.get("category_column")
    project_col = mapping.get("project_column")

    # Cache categories and projects for fast resolution
    cat_name_map = {c.name.strip().lower(): c.id for c in db.query(Category).all()}
    proj_name_map = {p.name.strip().lower(): p.id for p in db.query(Project).all()}

    # Determine Transaction Type (Income vs Expense)
    tx_type = TransactionType.Expense if (payload.transaction_type and payload.transaction_type.lower() == "expense") else TransactionType.Income

    # Fixed amount override
    fixed_amt = None
    if payload.fixed_amount is not None and float(payload.fixed_amount) > 0:
        fixed_amt = float(payload.fixed_amount)
    elif payload.amount_override is not None and float(payload.amount_override) > 0:
        fixed_amt = float(payload.amount_override)

    # Build multi-signal duplicate detector
    detector = build_duplicate_detector(db)

    # Get max transaction number for current year
    year = datetime.now().year
    prefix_pattern = f"YNV-{year}-"
    existing_txs = db.query(Ledger.transaction_number).filter(Ledger.transaction_number.startswith(prefix_pattern)).all()
    max_num = 0
    for (num_str,) in existing_txs:
        try:
            n = int(num_str.split("-")[-1])
            if n > max_num:
                max_num = n
        except Exception:
            pass

    imported_count = 0
    skipped_count = 0
    total_imported_amount = 0.0

    core_cols = {amount_col, name_col, ref_col, date_col, method_col, category_col, project_col, "_row_id", "_is_duplicate", "_duplicate_reason", "_parsed_amount"}

    for idx, r in enumerate(payload.rows):
        amt = 0.0
        if fixed_amt is not None:
            amt = fixed_amt
        elif amount_col and amount_col != "__FIXED__":
            amt = clean_amount(r.get(amount_col))
        
        if amt <= 0 and r.get("_parsed_amount"):
            amt = clean_amount(r.get("_parsed_amount"))

        if amt <= 0:
            skipped_count += 1
            continue

        raw_ref = str(r.get(ref_col) or "").strip() if ref_col else ""

        # Multi-signal duplicate check
        is_dup, dup_reason = detect_row_duplicate(r, amount_col, name_col, ref_col, fixed_amt, detector, idx + 1)
        if payload.skip_duplicates and (r.get("_is_duplicate") or is_dup):
            skipped_count += 1
            continue

        payer_name = str(r.get(name_col) or "").strip() if name_col else ""
        
        # Parse date
        tx_date = date.today()
        if date_col and r.get(date_col):
            d_str = str(r.get(date_col)).strip()
            for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d"):
                try:
                    tx_date = datetime.strptime(d_str.split()[0], fmt).date()
                    break
                except ValueError:
                    pass

        # Payment method
        p_method = str(r.get(method_col) or "").strip() if method_col else ""
        if not p_method:
            p_method = payload.default_payment_method or "Scanner / UPI"

        # Resolve Category
        row_cat_id = payload.category_id or None
        if category_col and r.get(category_col):
            val = str(r.get(category_col)).strip()
            if val.lower() in cat_name_map:
                row_cat_id = cat_name_map[val.lower()]
            elif val:
                new_cat = Category(
                    id=str(uuid.uuid4()),
                    name=val,
                    type=tx_type.value,
                    color="#10B981" if tx_type == TransactionType.Income else "#EF4444",
                    icon="tag",
                    is_active=True
                )
                db.add(new_cat)
                db.flush()
                cat_name_map[val.lower()] = new_cat.id
                row_cat_id = new_cat.id

        # Resolve Project
        row_proj_id = payload.project_id or None
        if project_col and r.get(project_col):
            val = str(r.get(project_col)).strip()
            if val.lower() in proj_name_map:
                row_proj_id = proj_name_map[val.lower()]
            elif val:
                clean_code = (re.sub(r'[^A-Za-z0-9]', '', val)[:4].upper() or "PRJ") + f"-{datetime.now().strftime('%m%d%H%M')}"
                new_proj = Project(
                    id=str(uuid.uuid4()),
                    name=val,
                    project_code=clean_code,
                    description="Auto-created from import",
                    status="Active"
                )
                db.add(new_proj)
                db.flush()
                proj_name_map[val.lower()] = new_proj.id
                row_proj_id = new_proj.id

        # Custom fields (any extra columns that were in the spreadsheet)
        custom_fields = {}
        for k, v in r.items():
            if k not in core_cols and str(v).strip():
                custom_fields[k] = str(v).strip()

        # Build informative description
        desc_parts = []
        if payer_name:
            desc_parts.append(payer_name)
        else:
            default_label = "Payment Collection" if tx_type == TransactionType.Income else "Expenditure Disbursement"
            desc_parts.append(default_label)

        for k, v in custom_fields.items():
            desc_parts.append(f"{k}: {v}")

        final_desc = " | ".join(desc_parts)

        # Generate sequential transaction number
        max_num += 1
        tx_number = f"YNV-{year}-{max_num:06d}"

        new_entry = Ledger(
            id=str(uuid.uuid4()),
            transaction_number=tx_number,
            type=tx_type,
            status=TransactionStatus.COMPLETED,
            amount=amt,
            category_id=row_cat_id,
            project_id=row_proj_id,
            description=final_desc,
            payment_method=p_method,
            reference_number=raw_ref if raw_ref else None,
            custom_metadata=json.dumps(custom_fields, ensure_ascii=False) if custom_fields else None,
            entered_by=str(current_user.id),
            transaction_date=tx_date
        )

        db.add(new_entry)
        if raw_ref and not ref_is_generic:
            existing_refs.add(raw_ref.lower())

        imported_count += 1
        total_imported_amount += amt

    if imported_count > 0:
        db.commit()
        AuditService.log_action(
            db=db,
            user_id=current_user.uuid,
            action=f"BULK_IMPORT_{tx_type.value.upper()}",
            module="Ledger",
            entity_type="Batch",
            entity_id=f"BATCH-{int(datetime.now().timestamp())}",
            old_values=None,
            new_values={
                "type": tx_type.value,
                "imported_count": imported_count,
                "total_amount": total_imported_amount,
                "skipped_count": skipped_count
            }
        )
    elif skipped_count > 0:
        raise HTTPException(
            status_code=400,
            detail=f"No transactions were imported. Checked {len(payload.rows)} row(s): all were skipped because amount was ₹0 or duplicate reference numbers were detected."
        )

    return {
        "message": f"Successfully imported {imported_count} {tx_type.value.lower()} transaction(s) totaling ₹{total_imported_amount:,.2f}",
        "imported_count": imported_count,
        "skipped_count": skipped_count,
        "total_amount": round(total_imported_amount, 2),
        "type": tx_type.value
    }

